"""Clean the school's club and student rosters into import-ready private files.

Read-only on the source workbooks. Corrections (club-name aliases, major
aliases, student-number fixes) come from an explicit JSON file so this script
never embeds school data. Any failed check aborts before anything is written.

Outputs (mode 0600) in --out: clubs.json, students.json, report.md.
"""

import argparse
import collections
import json
import os
from pathlib import Path
import re

from openpyxl import load_workbook

CLUB_COLUMNS = {
    "category": "一级社团",
    "name": "社团信息",
    "teacher": "指导老师姓名（多个则;隔开）",
    "president": "社长姓名",
    "summary": "社团简介",
    "capacity": "总名额",
    "reserved": "预留名额",
    "remaining": "剩余名额",
    "restricted": "是否有专业限制（有则1，无则0）",
    "majors": "allowed_majors（逗号分隔）",
}
STUDENT_COLUMNS = {
    "username": "学号",
    "name": "姓名",
    "class_name": "班级",
    "major": "专业",
    "college": "学院",
    "flag": "是否预留（1表示预留0表示非预留）",
    "club": "预留社团名称",
}


def text(value):
    return "" if value is None else str(value).strip()


def read_sheet(path, columns):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [text(v) for v in next(rows)]
    missing = [title for title in columns.values() if title not in header]
    if missing:
        raise SystemExit(f"{path.name} 缺少表头：{missing}")
    index = {key: header.index(title) for key, title in columns.items()}
    extra = [i for i, title in enumerate(header) if not title]
    result = []
    for number, row in enumerate(rows, 2):
        if not any(v is not None for v in row):
            continue
        item = {key: row[i] if i < len(row) else None for key, i in index.items()}
        item["row"] = number
        item["notes"] = [text(row[i]) for i in extra if i < len(row) and text(row[i])]
        result.append(item)
    return result


def main(args):
    fixes = json.loads(Path(args.corrections).read_text())
    club_aliases = fixes.get("club_aliases", {})
    major_aliases = fixes.get("major_aliases", {})
    username_fixes = fixes.get("username_fixes", {})
    description_overrides = fixes.get("description_overrides", {})
    errors, report = [], collections.defaultdict(list)

    clubs, category = [], ""
    for c in read_sheet(Path(args.clubs), CLUB_COLUMNS):
        category = text(c["category"]) or category
        name = text(c["name"])
        capacity, reserved, remaining = (int(c[k] or 0) for k in ("capacity", "reserved", "remaining"))
        if capacity - reserved != remaining:
            errors.append(f"社团表第{c['row']}行 {name}：总名额-预留≠剩余")
        majors = []
        if text(c["restricted"]) == "1":
            for raw in filter(None, (m.strip() for m in re.split(r"[,，、;；]", text(c["majors"])))):
                mapped = major_aliases.get(raw, [raw])
                if mapped != [raw]:
                    report["专业限制映射"].append(f"{name}：{raw} → {'、'.join(mapped)}")
                majors += [m for m in mapped if m not in majors]
            if not majors:
                errors.append(f"社团表第{c['row']}行 {name}：标记限专业但没有专业")
        summary = text(c["summary"])
        teacher, president = text(c["teacher"]), text(c["president"])
        if name in description_overrides:
            report["简介替换"].append(f"{name}：原简介「{summary}」→「{description_overrides[name]}」")
            summary = description_overrides[name]
        credits = "；".join(
            part for part in (teacher and f"指导老师：{teacher}", president and f"社长：{president}") if part
        )
        if not president:
            report["无社长"].append(name)
        clubs.append(
            {
                "name": name,
                "category": category,
                "description": "\n\n".join(part for part in (summary, credits) if part),
                "capacity": capacity,
                "allowed_majors": majors,
                "reserved": reserved,
                "remaining": remaining,
            }
        )
    names = [c["name"] for c in clubs]
    for name, count in collections.Counter(names).items():
        if count > 1:
            errors.append(f"社团名重复：{name}")
    for c in clubs:
        if not 0 < len(c["name"]) <= 64 or len(c["category"]) > 64:
            errors.append(f"{c['name']}：名称或分类长度超出限制")
    by_name = {c["name"]: c for c in clubs}

    students = []
    for s in read_sheet(Path(args.students), STUDENT_COLUMNS):
        username, name = text(s["username"]), text(s["name"])
        if username in username_fixes:
            report["学号更正"].append(f"第{s['row']}行 {name}：{username} → {username_fixes[username]}")
            username = username_fixes[username]
        club = text(s["club"])
        if club in club_aliases:
            report["预留社团名修正"].append(f"{club} → {club_aliases[club]}")
            club = club_aliases[club]
        flag = text(s["flag"])
        if flag == "":
            report["预留标记为空（按0处理）"].append(f"第{s['row']}行 {username} {name}")
            flag = "0"
        if (flag == "1") != bool(club):
            errors.append(f"学生表第{s['row']}行 {username}：预留标记与预留社团不一致")
        if not re.fullmatch(r"\d{10}", username):
            errors.append(f"学生表第{s['row']}行 学号格式异常：{username}")
        if not name or len(name.encode()) > 72:
            errors.append(f"学生表第{s['row']}行 {username}：姓名为空或过长")
        class_name, major = text(s["class_name"]), text(s["major"])
        year = re.search(r"(\d{2})\d{2}$", class_name)
        if year and username[2:4] != year.group(1):
            report["学号年份与班级年级不一致（照录）"].append(f"{username} {name} {class_name}")
        if not major:
            report["无专业（留空）"].append(f"{username} {name} {class_name} {text(s['college'])}")
        if s["notes"]:
            report["备注列（不导入）"].append(f"{username} {name}：{'；'.join(s['notes'])}")
        if club:
            target = by_name.get(club)
            if not target:
                errors.append(f"学生表第{s['row']}行 {username}：预留社团不存在：{club}")
            elif target["allowed_majors"] and major not in target["allowed_majors"]:
                errors.append(f"学生表第{s['row']}行 {username}：专业{major}不符合{club}的限制")
        students.append(
            {"username": username, "name": name, "major": major, "class_name": class_name, "reserved_club": club}
        )
    for username, count in collections.Counter(s["username"] for s in students).items():
        if count > 1:
            errors.append(f"学号重复：{username}")
    reserved = collections.Counter(s["reserved_club"] for s in students if s["reserved_club"])
    for c in clubs:
        if reserved[c["name"]] != c["reserved"]:
            errors.append(f"{c['name']}：名单预留{reserved[c['name']]}人，社团表预留名额{c['reserved']}")
        if reserved[c["name"]] > c["capacity"]:
            errors.append(f"{c['name']}：预留人数超过总名额")
        if c["capacity"] and not c["remaining"]:
            report["预留已占满全部名额"].append(c["name"])

    if errors:
        raise SystemExit("核对未通过，未写出任何文件：\n" + "\n".join(errors))

    out = Path(args.out)
    out.mkdir(mode=0o700, parents=True, exist_ok=True)
    summary = {
        "clubs": len(clubs),
        "categories": len({c["category"] for c in clubs}),
        "capacity": sum(c["capacity"] for c in clubs),
        "students": len(students),
        "reserved": sum(reserved.values()),
        "open_seats": sum(c["capacity"] for c in clubs) - sum(reserved.values()),
        "unreserved_students": sum(1 for s in students if not s["reserved_club"]),
    }
    lines = ["# 学校数据核对报告", "", "## 汇总", ""]
    lines += [f"- {k}: {v}" for k, v in summary.items()]
    for title, items in report.items():
        counted = collections.Counter(items)
        lines += ["", f"## {title}（{len(items)}）", ""]
        lines += [f"- {item}" + (f" ×{n}" if n > 1 else "") for item, n in counted.items()]
    files = {
        "clubs.json": json.dumps(clubs, ensure_ascii=False, indent=1),
        "students.json": json.dumps(students, ensure_ascii=False, indent=1),
        "report.md": "\n".join(lines) + "\n",
    }
    for filename, content in files.items():
        path = out / filename
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        with os.fdopen(fd, "w") as f:
            f.write(content)
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--clubs", required=True)
    p.add_argument("--students", required=True)
    p.add_argument("--corrections", required=True)
    p.add_argument("--out", required=True)
    main(p.parse_args())
