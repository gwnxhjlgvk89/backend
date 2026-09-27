import asyncio
from io import BytesIO
import zipfile
import base64
import hashlib

from cryptography.fernet import Fernet

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from club.config import settings
from club.errors import fail
from club.models import Activity, Audit, Club, Enrollment, Job, User, utcnow
from club.schemas import StudentInput
from club.services import mutable_activity


def upload_path(job_id):
    return settings().storage_dir / "private" / f"{job_id}.encrypted"


def upload_cipher():
    return Fernet(base64.urlsafe_b64encode(hashlib.sha256(
        (settings().jwt_secret + ":import-upload").encode()).digest()))


def persist_upload(job_id, content):
    path = upload_path(job_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as f:
        path.chmod(0o600)
        f.write(upload_cipher().encrypt(content))


def parse_students(content):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                fail(413, "FILE_TOO_LARGE", "解压后文件过大")
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows = book.active.iter_rows(values_only=True)
        header = next(rows)
        expected = ("username", "name", "major", "class_name", "password", "reserved_club_id")
        if tuple(header[:6]) != expected:
            fail(422, "INVALID_COLUMNS", "表头须为 username,name,major,class_name,password,reserved_club_id")
        result, seen = [], set()
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            if len(result) >= 5000:
                fail(413, "TOO_MANY_ROWS", "一次最多导入5000行")
            data = dict(zip(expected, row))
            for field in expected[:5]:
                data[field] = str(data.get(field) or "").strip()
            data["reserved_club_id"] = data.get("reserved_club_id") or None
            validated = StudentInput.model_validate(data)
            if not validated.password or validated.username in seen:
                fail(422, "INVALID_ROW", f"第{number}行密码缺失或账号重复")
            seen.add(validated.username)
            result.append(validated.model_dump())
        book.close()
        if not result:
            fail(422, "EMPTY_FILE", "文件没有有效数据")
        return result
    except (zipfile.BadZipFile, ValueError, StopIteration, KeyError):
        fail(422, "INVALID_WORKBOOK", "文件格式或单元格内容不正确")


def write_export(path, rows):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("报名结果")
    sheet.append(["学号", "姓名", "专业", "班级", "社团", "来源"])
    for row in rows:
        # Prevent spreadsheet formula injection.
        sheet.append([("'" + str(v)) if str(v or "").startswith(("=", "+", "-", "@")) else v for v in row])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    path.chmod(0o600)


async def run_one(factory):
    async with factory() as db, db.begin():
        a = await db.get(Activity, 1, with_for_update={"read": True})
        if a and a.state == "published" and a.starts_at <= utcnow() < a.ends_at:
            return
        job = await db.scalar(select(Job).where(Job.state == "queued").order_by(Job.created_at).limit(1).with_for_update(skip_locked=True))
        if not job:
            return
        job_id, kind, payload = job.id, job.kind, job.payload
        job.state = "running"
    try:
        if kind == "import_preview":
            from club.security import hash_password
            path = upload_path(job_id)
            content = await asyncio.to_thread(path.read_bytes)
            rows = await asyncio.to_thread(parse_students, upload_cipher().decrypt(content))
            for row in rows:
                row["password_hash"] = await hash_password(row.pop("password"))
            async with factory() as db, db.begin():
                await mutable_activity(db)
                existing = set(await db.scalars(select(User.username).where(User.username.in_([r["username"] for r in rows]))))
                if existing:
                    fail(409, "DUPLICATE_STUDENTS", "文件包含已有账号，请移除后重新导入")
                clubs = {c.id: c for c in await db.scalars(select(Club))}
                reservations = {}
                for row in rows:
                    cid = row["reserved_club_id"]
                    if cid:
                        c = clubs.get(cid)
                        if not c or not c.active or (c.allowed_majors and row["major"] not in c.allowed_majors):
                            fail(409, "INVALID_RESERVATION", "预报名社团不存在、未开放或不符合专业要求")
                        reservations[cid] = reservations.get(cid, 0) + 1
                        if reservations[cid] > c.remaining:
                            fail(409, "RESERVATION_FULL", "预报名人数超过剩余名额")
                j = await db.get(Job, job_id)
                j.state, j.payload = "preview", {"rows": rows}
                j.result = {"rows": len(rows), "preview": [{k: r[k] for k in ("username", "name", "major", "class_name")} for r in rows[:20]]}
            path.unlink(missing_ok=True)
        elif kind == "import_students":
            from club.admin import save_student
            async with factory() as db, db.begin():
                await mutable_activity(db)
                for row in payload["rows"]:
                    body = StudentInput.model_validate({k: v for k, v in row.items() if k != "password_hash"})
                    await save_student(db, body, row["password_hash"])
                db.add(Audit(actor_id=job.actor_id, action="import.apply", target=job_id, summary={"rows": len(payload["rows"])}))
                result = {"rows": len(payload["rows"])}
                j = await db.get(Job, job_id)
                j.state, j.result, j.payload = "done", result, {}
        elif kind == "export_students":
            async with factory() as db, db.begin():
                await mutable_activity(db)
                rows = (await db.execute(select(User.username, User.name, User.major, User.class_name, Club.name, Enrollment.source)
                    .outerjoin(Enrollment, Enrollment.student_id == User.id).outerjoin(Club, Club.id == Enrollment.club_id)
                    .where(User.role == "student").order_by(User.id))).all()
            path = settings().storage_dir / "private" / f"{job_id}.xlsx"
            await asyncio.to_thread(write_export, path, rows)
            async with factory() as db, db.begin():
                j = await db.get(Job, job_id)
                j.state, j.result = "done", {"rows": len(rows), "download": f"/admin/jobs/{job_id}/download"}
        else:
            raise ValueError("unknown job kind")
    except Exception as exc:
        from club.errors import DomainError
        upload_path(job_id).unlink(missing_ok=True)
        async with factory() as db, db.begin():
            j = await db.get(Job, job_id)
            message = exc.message if isinstance(exc, DomainError) else "任务失败，数据未部分导入，请检查文件或活动状态"
            j.state, j.result, j.payload = "failed", {"message": message, "type": type(exc).__name__}, {}
